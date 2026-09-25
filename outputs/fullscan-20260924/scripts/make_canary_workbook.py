"""Tạo workbook canary 1 hotel (Mercure Vũng Tàu) lấy đúng dòng trong link_hotel_data_expanded.xlsx. Claude, 2026-09-25.

Canary = chạy 1 trang đầu phiên Chrome mới để xem Booking đang ở trạng thái BẬT (có dòng "Chỉ dành cho 1 khách") hay TẮT; xem README_fullscan.md.
Không sửa workbook nguồn; ghi run/canary_mercure.xlsx (gitignore theo run/).
Chạy:  python scripts/make_canary_workbook.py   (venv backend, có openpyxl)
"""
import hashlib
from pathlib import Path

import openpyxl

SRC = Path("D:/MSE/CAPSTONE/link_hotel_data_expanded.xlsx")
OUT = Path(__file__).resolve().parents[1] / "run" / "canary_mercure.xlsx"
SLUG = "mercure-vung-tau-vietnam"

wb = openpyxl.load_workbook(SRC)
out = openpyxl.Workbook()
out.remove(out.active)
found = 0
for ws in wb.worksheets:
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r and len(r) > 1 and r[1] and f"/hotel/vn/{SLUG}." in str(r[1])]
    if rows:
        o = out.create_sheet(ws.title)
        o.append([c.value for c in ws[1]])
        for r in rows:
            o.append(list(r))
            found += 1
assert found == 1, f"cần đúng 1 dòng {SLUG}, thấy {found}"
OUT.parent.mkdir(parents=True, exist_ok=True)
out.save(OUT)
print(f"đã ghi {OUT} ({found} hotel, sheet {out.sheetnames}); SHA-256 {hashlib.sha256(OUT.read_bytes()).hexdigest()}")
