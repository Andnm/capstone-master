"""Cohort manifest + source manifest + ownership workbook reader - dung file tam that.

Khong mock openpyxl: cac loi that da gap (o cong thuc chua tinh, sheet sai ten, thieu cot) chi xuat
hien khi doc file that, nen fixture o day cung ghi file .xlsx/.json that vao tmp_path.
"""
import datetime as dt
import json

import openpyxl
import pytest

from app.warehouse.cohort_manifest import VALID_CITIES, load_cohort_manifest
from app.warehouse.errors import ManifestError
from app.warehouse.ownership_manifest import (
    OwnershipManifest,
    OwnershipRow,
    build_ownership_rows,
    load_ownership_manifest,
    read_plan_rows,
    write_ownership_manifest,
)
from app.warehouse.source_manifest import (
    extract_schema_text,
    load_source_manifest,
    verify_dump_checksum,
)

# ======================================================================================
# Cohort manifest
# ======================================================================================
def write_cohort(path, sheets):
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for city, rows in sheets.items():
        sheet = workbook.create_sheet(city)
        sheet.append(["Tên khách sạn", "Link"])
        for name, slug in rows:
            sheet.append([name, f"https://www.booking.com/hotel/vn/{slug}.vi.html?aid=1"])
    workbook.save(path)
    return path


def test_cohort_doc_duoc_hotel_id_va_city(tmp_path):
    path = write_cohort(tmp_path / "cohort.xlsx", {
        "Đà Lạt": [("A", "mac-dalat"), ("B", "lumina-dalat-premium")],
        "Phú Quốc": [("C", "roma")],
    })
    manifest = load_cohort_manifest(path, require_all_cities=False)
    assert manifest.size == 3
    assert manifest.city_of("mac-dalat") == "Đà Lạt"
    assert manifest.city_of("roma") == "Phú Quốc"
    assert manifest.contains("lumina-dalat-premium")


def test_cohort_hotel_ngoai_danh_sach_tra_city_none(tmp_path):
    """Hotel ngoai cohort -> city NULL o warehouse (muc 6 rule 5), khong fallback city cua nguon."""
    path = write_cohort(tmp_path / "cohort.xlsx", {"Đà Lạt": [("A", "mac-dalat")]})
    manifest = load_cohort_manifest(path, require_all_cities=False)
    assert manifest.city_of("sen") is None
    assert not manifest.contains("sen")


def test_cohort_hash_khong_doi_khi_luu_lai_file(tmp_path):
    """Hash tinh tu NOI DUNG, khong tu bytes .xlsx (Excel doi bytes moi lan luu)."""
    sheets = {"Đà Lạt": [("A", "mac-dalat")], "Hà Nội": [("B", "roma")]}
    first = load_cohort_manifest(write_cohort(tmp_path / "a.xlsx", sheets), require_all_cities=False)
    second = load_cohort_manifest(write_cohort(tmp_path / "b.xlsx", sheets), require_all_cities=False)
    assert first.manifest_sha256 == second.manifest_sha256


def test_cohort_sheet_khong_phai_thanh_pho_thi_fail(tmp_path):
    path = write_cohort(tmp_path / "cohort.xlsx", {"Đà Nẵng": [("A", "x")]})
    with pytest.raises(ManifestError, match="khong phai thanh pho"):
        load_cohort_manifest(path)


def test_cohort_link_khong_parse_duoc_slug_thi_fail(tmp_path):
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    sheet = workbook.create_sheet("Đà Lạt")
    sheet.append(["Tên khách sạn", "Link"])
    sheet.append(["A", "https://example.com/khong-phai-booking"])
    workbook.save(tmp_path / "cohort.xlsx")
    with pytest.raises(ManifestError, match="khong parse duoc slug"):
        load_cohort_manifest(tmp_path / "cohort.xlsx", require_all_cities=False)


def test_cohort_hotel_trung_o_2_sheet_thi_fail(tmp_path):
    path = write_cohort(tmp_path / "cohort.xlsx", {
        "Đà Lạt": [("A", "roma")], "Phú Quốc": [("B", "roma")],
    })
    with pytest.raises(ManifestError, match="nhieu sheet"):
        load_cohort_manifest(path, require_all_cities=False)


# ======================================================================================
# Ownership workbook reader
# ======================================================================================
def write_plan(path, *, sheet, headers, rows):
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    worksheet = workbook.create_sheet(sheet)
    worksheet.append(["tieu de"])
    worksheet.append([])
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    workbook.save(path)
    return path


VPS_HEADERS = ["Crawl date", "V1 weekday", "V2 weekend", "V3 seasonal", "V4 medium", "V5 long"]


def test_doc_plan_vps(tmp_path):
    path = write_plan(tmp_path / "vps.xlsx", sheet="VPS_CRAWL_PLAN", headers=VPS_HEADERS, rows=[
        [dt.datetime(2026, 8, 24), dt.datetime(2026, 8, 27), dt.datetime(2026, 8, 30),
         dt.datetime(2026, 9, 3), dt.datetime(2026, 9, 16), dt.datetime(2027, 1, 24)],
    ])
    rows = read_plan_rows("vps", path)
    assert len(rows) == 5
    assert {row.schedule_slot for row in rows} == {"V1", "V2", "V3", "V4", "V5"}
    assert all(row.crawl_date == dt.date(2026, 8, 24) for row in rows)


def test_o_trong_duoc_bo_qua_khong_phai_loi(tmp_path):
    """AN2/AN3 cua may phu bi xoa tu 03/09 - o trong la hop le, khong phai thieu du lieu."""
    path = write_plan(tmp_path / "vps.xlsx", sheet="VPS_CRAWL_PLAN", headers=VPS_HEADERS, rows=[
        [dt.datetime(2026, 8, 24), dt.datetime(2026, 8, 27), None, None,
         dt.datetime(2026, 9, 16), dt.datetime(2027, 1, 24)],
    ])
    assert len(read_plan_rows("vps", path)) == 3


def test_o_cong_thuc_chua_tinh_thi_fail_ro_rang(tmp_path):
    """Dung ca that cua aux_local_crawl_sampling_master.xlsx: 1.080 o la cong thuc chua cached."""
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    sheet = workbook.create_sheet("VPS_CRAWL_PLAN")
    sheet.append(["tieu de"])
    sheet.append([])
    sheet.append(VPS_HEADERS)
    sheet.append([dt.datetime(2026, 8, 24), "='SLOT_RULES'!$D$4+1", None, None, None, None])
    workbook.save(tmp_path / "formula.xlsx")
    with pytest.raises(ManifestError, match="CONG THUC"):
        read_plan_rows("vps", tmp_path / "formula.xlsx")


def test_thieu_cot_slot_thi_fail(tmp_path):
    path = write_plan(tmp_path / "vps.xlsx", sheet="VPS_CRAWL_PLAN",
                      headers=["Crawl date", "V1 weekday"], rows=[[dt.datetime(2026, 8, 24), None]])
    with pytest.raises(ManifestError, match="thieu cot slot"):
        read_plan_rows("vps", path)


def test_source_code_chua_khai_bao_layout_thi_fail(tmp_path):
    with pytest.raises(ManifestError, match="chua khai bao layout"):
        read_plan_rows("nguon_la", tmp_path / "x.xlsx")


def test_ghi_va_doc_lai_ownership_manifest(tmp_path):
    plan = write_plan(tmp_path / "vps.xlsx", sheet="VPS_CRAWL_PLAN", headers=VPS_HEADERS, rows=[
        [dt.datetime(2026, 8, 24), dt.datetime(2026, 8, 27), dt.datetime(2026, 8, 30),
         dt.datetime(2026, 9, 3), dt.datetime(2026, 9, 16), dt.datetime(2027, 1, 24)],
    ])
    rows = build_ownership_rows({"vps": plan})
    out = tmp_path / "ownership.json"
    digest = write_ownership_manifest(rows, out)
    manifest = load_ownership_manifest(out)
    assert manifest.manifest_sha256 == digest
    assert manifest.lookup(dt.date(2026, 8, 24), dt.date(2026, 8, 27)).schedule_slot == "V1"
    assert manifest.lookup(dt.date(2026, 8, 24), dt.date(2099, 1, 1)) is None


def test_cohort_that_du_5_thanh_pho_va_354_hotel():
    """Kiem tra chinh file cohort THAT dang dung, khong phai fixture - neu ai do xoa mot sheet thi
    scope dataset am tham co lai ma khong ai biet (GPT review 06 MINOR 4)."""
    from pathlib import Path

    path = Path(__file__).resolve().parents[3] / "link_hotel_data_expanded.xlsx"
    if not path.exists():
        pytest.skip(f"khong tim thay cohort that tai {path}")
    manifest = load_cohort_manifest(path, require_all_cities=False)
    assert manifest.size == 354, f"cohort v2 phai la 354 hotel, dang la {manifest.size}"
    assert set(manifest.hotel_city.values()) == set(VALID_CITIES)


def test_thieu_sheet_thanh_pho_thi_fail_mac_dinh(tmp_path):
    """Xoa nham 1 sheet khong duoc lam scope am tham co tu 5 thanh pho xuong 4."""
    path = write_cohort(tmp_path / "cohort.xlsx", {"Đà Lạt": [("A", "mac-dalat")]})
    with pytest.raises(ManifestError, match="THIEU sheet thanh pho"):
        load_cohort_manifest(path)


def test_cohort_mapping_khong_sua_duoc_sau_khi_load(tmp_path):
    path = write_cohort(tmp_path / "cohort.xlsx", {"Đà Lạt": [("A", "mac-dalat")]})
    manifest = load_cohort_manifest(path, require_all_cities=False)
    with pytest.raises(TypeError):
        manifest.hotel_city["moi"] = "Hà Nội"


# --- Validation fail-closed cua row/envelope (GPT review 06 MAJOR 1) ---------------------
def test_owner_source_khong_khai_bao_layout_thi_fail():
    with pytest.raises(ManifestError, match="WORKBOOK_LAYOUTS"):
        OwnershipManifest.from_rows([
            OwnershipRow("vpz", dt.date(2026, 9, 14), "V1", dt.date(2026, 9, 27)),
        ])


def test_slot_khong_thuoc_nguon_do_thi_fail():
    """Slot 'N1' la cua local_primary; VPS khong duoc dung."""
    with pytest.raises(ManifestError, match="khong thuoc slot cua nguon"):
        OwnershipManifest.from_rows([
            OwnershipRow("vps", dt.date(2026, 9, 14), "N1", dt.date(2026, 9, 27)),
        ])


def test_slot_rong_thi_fail():
    with pytest.raises(ManifestError, match="schedule_slot rong"):
        OwnershipManifest.from_rows([
            OwnershipRow("vps", dt.date(2026, 9, 14), "  ", dt.date(2026, 9, 27)),
        ])


@pytest.mark.parametrize("field,bad", [("row_count", 99), ("sources", ["local_primary"])])
def test_envelope_khong_khop_thi_fail(tmp_path, field, bad):
    plan = write_plan(tmp_path / "vps.xlsx", sheet="VPS_CRAWL_PLAN", headers=VPS_HEADERS, rows=[
        [dt.datetime(2026, 8, 24), dt.datetime(2026, 8, 27), None, None, None, None],
    ])
    out = tmp_path / "ownership.json"
    write_ownership_manifest(build_ownership_rows({"vps": plan}), out)
    document = json.loads(out.read_text(encoding="utf-8"))
    document[field] = bad
    out.write_text(json.dumps(document), encoding="utf-8")
    with pytest.raises(ManifestError, match=field):
        load_ownership_manifest(out)


def test_ghi_manifest_la_atomic_khong_de_lai_file_cut(tmp_path):
    plan = write_plan(tmp_path / "vps.xlsx", sheet="VPS_CRAWL_PLAN", headers=VPS_HEADERS, rows=[
        [dt.datetime(2026, 8, 24), dt.datetime(2026, 8, 27), None, None, None, None],
    ])
    rows = build_ownership_rows({"vps": plan})
    out = tmp_path / "sub" / "ownership.json"
    write_ownership_manifest(rows, out)
    assert out.exists()
    # Khong con file tam nao sot lai trong thu muc dich.
    assert [p.name for p in out.parent.iterdir()] == ["ownership.json"]


def _tampered_manifest(tmp_path, mutate):
    plan = write_plan(tmp_path / "vps.xlsx", sheet="VPS_CRAWL_PLAN", headers=VPS_HEADERS, rows=[
        [dt.datetime(2026, 8, 24), dt.datetime(2026, 8, 27), None, None, None, None],
    ])
    out = tmp_path / "ownership.json"
    write_ownership_manifest(build_ownership_rows({"vps": plan}), out)
    document = json.loads(out.read_text(encoding="utf-8"))
    mutate(document)
    out.write_text(json.dumps(document), encoding="utf-8")
    return out


def test_sua_slot_thanh_slot_la_bi_bat_o_validate_row(tmp_path):
    out = _tampered_manifest(tmp_path, lambda d: d["rows"][0].update(schedule_slot="V9"))
    with pytest.raises(ManifestError, match="khong thuoc slot cua nguon"):
        load_ownership_manifest(out)


def test_sua_ngay_thanh_ngay_HOP_LE_khac_van_bi_bat_bang_hash(tmp_path):
    """Row van hop le ve moi mat -> chi con hash bat duoc. Day la test rieng cho lop hash,
    vi validate_rows() moi them co the che mat no neu chi tamper slot."""
    out = _tampered_manifest(tmp_path, lambda d: d["rows"][0].update(checkin_date="2026-08-28"))
    with pytest.raises(ManifestError, match="sua tay"):
        load_ownership_manifest(out)


# ======================================================================================
# Source manifest
# ======================================================================================
_DUMP = """/*!40101 SET @saved=@@x */;
DROP TABLE IF EXISTS `hotels`;
CREATE TABLE `hotels` (
  `hotel_id` varchar(255) NOT NULL,
  PRIMARY KEY (`hotel_id`)
) ENGINE=InnoDB AUTO_INCREMENT=42 DEFAULT CHARSET=utf8mb4;
INSERT INTO `hotels` VALUES ('a');
"""


def write_source_manifest(tmp_path, **overrides):
    dump = tmp_path / "dump.sql"
    dump.write_text(_DUMP, encoding="utf-8")
    from app.warehouse.hashing import file_sha256

    entry = {
        "source_code": "vps",
        "source_priority": 0,
        "dump_path": "dump.sql",
        "dump_sha256": file_sha256(dump),
        "dump_taken_at": "2026-09-15T15:49:12Z",
        "schema_sha256": "f" * 64,
        "source_version_json": {"scraper_version": "2.3.0"},
    }
    entry.update(overrides)
    path = tmp_path / "sources.json"
    path.write_text(json.dumps({"manifest_version": 1, "sources": [entry]}), encoding="utf-8")
    return path


def test_load_source_manifest_va_verify_dump(tmp_path):
    manifest = load_source_manifest(write_source_manifest(tmp_path), base_dir=tmp_path)
    assert [entry.source_code for entry in manifest.sources] == ["vps"]
    assert verify_dump_checksum(manifest.sources[0]) == manifest.sources[0].dump_sha256


def test_dump_bi_sua_sau_khi_manifest_tao_thi_fail(tmp_path):
    path = write_source_manifest(tmp_path)
    (tmp_path / "dump.sql").write_text(_DUMP + "-- them mot dong\n", encoding="utf-8")
    manifest = load_source_manifest(path, base_dir=tmp_path)
    with pytest.raises(ManifestError, match="dump_sha256 KHONG KHOP"):
        verify_dump_checksum(manifest.sources[0])


def test_source_code_sai_dinh_dang_thi_fail(tmp_path):
    path = write_source_manifest(tmp_path, source_code="VPS-1")
    with pytest.raises(ManifestError, match="source_code"):
        load_source_manifest(path, base_dir=tmp_path)


def test_source_priority_am_thi_fail(tmp_path):
    path = write_source_manifest(tmp_path, source_priority=-1)
    with pytest.raises(ManifestError, match="source_priority"):
        load_source_manifest(path, base_dir=tmp_path)


def test_sha256_khong_dung_dinh_dang_thi_fail(tmp_path):
    path = write_source_manifest(tmp_path, schema_sha256="qua-ngan")
    with pytest.raises(ManifestError, match="SHA-256"):
        load_source_manifest(path, base_dir=tmp_path)


def test_priority_trung_giua_2_nguon_thi_fail(tmp_path):
    dump = tmp_path / "dump.sql"
    dump.write_text(_DUMP, encoding="utf-8")
    from app.warehouse.hashing import file_sha256

    base = {
        "source_priority": 0, "dump_path": "dump.sql", "dump_sha256": file_sha256(dump),
        "dump_taken_at": "2026-09-15T15:49:12Z", "schema_sha256": "f" * 64,
        "source_version_json": {},
    }
    path = tmp_path / "sources.json"
    path.write_text(json.dumps({"manifest_version": 1, "sources": [
        dict(base, source_code="vps"), dict(base, source_code="local_primary"),
    ]}), encoding="utf-8")
    with pytest.raises(ManifestError, match="source_priority=0 bi trung"):
        load_source_manifest(path, base_dir=tmp_path)


def test_extract_schema_bo_auto_increment():
    """AUTO_INCREMENT=N phu thuoc so dong, khong phan anh cau truc - phai bi bo truoc khi hash."""
    text = extract_schema_text(_DUMP)
    assert "AUTO_INCREMENT" not in text
    assert "CREATE TABLE `hotels`" in text
    assert "INSERT INTO" not in text
