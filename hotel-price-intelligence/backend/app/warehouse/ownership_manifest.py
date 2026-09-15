"""Ownership manifest: ai "so huu" mot (crawl_date, checkin_date) - muc 8.

Manifest la bang PHANG sinh tu N workbook lich, checksum truoc khi dung. Moi dong:
`(owner_source, crawl_date, schedule_slot, checkin_date)`. Khoa tra cuu la
`(crawl_date, checkin_date)` - DUY NHAT tren TOAN BO manifest, khong phai duy nhat trong tung nguon.

Vi sao khoa khong gom source_code (chot voi GPT, file 02/02b): neu tra cuu theo
`owner_source = source_code` truoc thi khong bao gio phat hien duoc `non_owner_duplicate`. Phai tra
dong truoc, roi moi SO `manifest.owner_source` voi source thuc te. Kiem chung tren du lieu that:
VPS ngay 14/09/2026 da cao 3 check-in thuoc slot cua may chinh - chi thuat toan nay nhan ra.

"hotel/cohort" trong khoa ownership o muc 8 nghia la lich ap cho TOAN COHORT, khong sinh 354 dong
cho moi hotel. Cohort membership la gate rieng o cap item (-> `protocol_deviation`).
"""
from __future__ import annotations

import datetime as dt
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from .errors import ManifestError, OwnershipConflictError
from .hashing import canonical_json, sha256_hex

# Layout workbook cua tung nguon. Them nguon moi = them 1 entry o day, KHONG sua logic resolve.
# Ten cot lay dung nhu header dong 3 cua tung file; slot la ten ngan dung trong etl_item_map
# (VARCHAR(20)).
WORKBOOK_LAYOUTS: dict[str, dict[str, Any]] = {
    "local_primary": {
        "sheet": "DAILY_CRAWL_PLAN",
        "header_row": 3,
        "crawl_date_header": "Crawl date",
        "slots": {
            "N1": "N1 gần", "N2": "N2 gần", "N3": "N3 gần", "N4": "N4 gần", "N5": "N5 gần",
            "N6": "N6 gần", "N7": "N7 gần", "N8": "N8 gần", "N9": "N9 gần",
            "F1": "F1 xa", "F2": "F2 xa", "F3": "F3 xa",
        },
    },
    "vps": {
        "sheet": "VPS_CRAWL_PLAN",
        "header_row": 3,
        "crawl_date_header": "Crawl date",
        "slots": {
            "V1": "V1 weekday", "V2": "V2 weekend", "V3": "V3 seasonal",
            "V4": "V4 medium", "V5": "V5 long",
        },
    },
    "local_aux": {
        "sheet": "AUX_CRAWL_PLAN",
        "header_row": 3,
        "crawl_date_header": "Crawl date",
        "slots": {
            "AN1": "AN1", "AN2": "AN2 (chỉ 02/09)", "AN3": "AN3 (chỉ 02/09)",
            "AN4": "AN4", "AN5": "AN5", "AN6": "AN6", "AN7": "AN7", "AN8": "AN8", "AN9": "AN9",
            "AF1": "AF1", "AF2": "AF2", "AF3": "AF3",
        },
    },
}

# Ly do loai o cap item - dung ENUM chuoi co dinh de report gom nhom duoc, khong viet tu do.
REASON_PRE_PROTOCOL_PILOT = "pre_protocol_pilot"
REASON_POST_PROTOCOL_WINDOW = "post_protocol_window"
REASON_DAY_NOT_IN_MANIFEST = "day_not_in_ownership_manifest"
REASON_OFF_PLAN_CHECKIN_OTHER_DAY = "off_plan_checkin_other_day"
REASON_OFF_PLAN_UNKNOWN = "off_plan_unknown"
REASON_NON_OWNER_DUPLICATE = "non_owner_duplicate"
REASON_HOTEL_OUTSIDE_COHORT = "hotel_outside_cohort_manifest"
REASON_RUN_NOT_IN_MANIFEST = "run_not_in_ownership_manifest"

MAX_SLOT_LENGTH = 20  # etl_item_map.schedule_slot VARCHAR(20)


@dataclass(frozen=True)
class OwnershipRow:
    owner_source: str
    crawl_date: dt.date
    schedule_slot: str
    checkin_date: dt.date

    def as_payload(self) -> dict[str, Any]:
        return {
            "owner_source": self.owner_source,
            "crawl_date": self.crawl_date.isoformat(),
            "schedule_slot": self.schedule_slot,
            "checkin_date": self.checkin_date.isoformat(),
        }


@dataclass(frozen=True)
class EligibilityFlags:
    include_eda_raw: bool
    include_eda_main: bool
    include_reference: bool
    include_training: bool

    def intersect(self, other: "EligibilityFlags") -> "EligibilityFlags":
        """AND tung co - item khong bao gio duoc eligibility cao hon run cha (GPT file 04, O4)."""
        return EligibilityFlags(
            include_eda_raw=self.include_eda_raw and other.include_eda_raw,
            include_eda_main=self.include_eda_main and other.include_eda_main,
            include_reference=self.include_reference and other.include_reference,
            include_training=self.include_training and other.include_training,
        )


FLAGS_FULL = EligibilityFlags(True, True, True, True)
FLAGS_RAW_AND_MAIN_ONLY = EligibilityFlags(True, True, False, False)
FLAGS_RAW_ONLY = EligibilityFlags(True, False, False, False)


@dataclass(frozen=True)
class ItemOwnership:
    ownership_status: str
    schedule_slot: str | None
    owner_source: str | None
    flags: EligibilityFlags
    exclusion_reason: str | None


@dataclass(frozen=True)
class RunOwnership:
    planned_crawl_date: dt.date | None
    has_plan: bool
    flags: EligibilityFlags
    exclusion_reason: str | None


# ======================================================================================
# Sinh manifest tu workbook
# ======================================================================================
def read_plan_rows(source_code: str, workbook_path: str | Path) -> list[OwnershipRow]:
    """Doc 1 workbook lich -> cac dong ownership cua dung nguon do.

    FAIL RO neu o slot la CONG THUC chua co cached value (truong hop that cua
    `aux_local_crawl_sampling_master.xlsx`: file do script openpyxl ghi ra, chua tung mo bang Excel
    nen `data_only=True` tra None cho toan bo 1.080 o). Khong tu implement lai SLOT_RULES trong ETL
    (yeu cau cua GPT file 02b) - phai materialize o tang workbook.
    """
    import openpyxl

    layout = WORKBOOK_LAYOUTS.get(source_code)
    if layout is None:
        raise ManifestError(
            f"chua khai bao layout workbook cho source_code={source_code!r} - them vao "
            f"WORKBOOK_LAYOUTS truoc khi dung nguon nay."
        )
    path = Path(workbook_path)
    if not path.exists():
        raise ManifestError(f"nguon {source_code!r}: khong tim thay workbook {path}")

    values_wb = openpyxl.load_workbook(path, data_only=True)
    formula_wb = openpyxl.load_workbook(path, data_only=False)
    try:
        sheet_name = layout["sheet"]
        if sheet_name not in values_wb.sheetnames:
            raise ManifestError(
                f"nguon {source_code!r}: workbook khong co sheet {sheet_name!r} "
                f"(co: {values_wb.sheetnames})."
            )
        values_sheet = values_wb[sheet_name]
        formula_sheet = formula_wb[sheet_name]
        header_row = layout["header_row"]
        header = [cell for cell in next(values_sheet.iter_rows(
            min_row=header_row, max_row=header_row, values_only=True))]
        column_index = _map_headers(source_code, header, layout)

        rows: list[OwnershipRow] = []
        for excel_row, row in enumerate(
            values_sheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            crawl_date = _as_date(row[column_index["__crawl_date__"]]) if row else None
            if crawl_date is None:
                continue
            for slot, col in column_index.items():
                if slot == "__crawl_date__":
                    continue
                value = row[col] if col < len(row) else None
                if value is None:
                    _require_not_unevaluated_formula(
                        source_code, formula_sheet.cell(row=excel_row, column=col + 1), slot, excel_row
                    )
                    continue
                checkin = _as_date(value)
                if checkin is None:
                    raise ManifestError(
                        f"nguon {source_code!r} dong {excel_row} slot {slot}: gia tri {value!r} "
                        f"khong phai ngay."
                    )
                rows.append(OwnershipRow(
                    owner_source=source_code, crawl_date=crawl_date,
                    schedule_slot=slot, checkin_date=checkin,
                ))
    finally:
        values_wb.close()
        formula_wb.close()

    if not rows:
        raise ManifestError(f"nguon {source_code!r}: khong doc duoc dong lich nao tu {path}.")
    return rows


def _map_headers(source_code: str, header: Sequence[Any], layout: Mapping[str, Any]) -> dict[str, int]:
    labels = {str(cell).strip(): index for index, cell in enumerate(header) if cell is not None}
    mapping: dict[str, int] = {}
    crawl_header = layout["crawl_date_header"]
    if crawl_header not in labels:
        raise ManifestError(
            f"nguon {source_code!r}: khong tim thay cot {crawl_header!r} trong header {list(labels)}."
        )
    mapping["__crawl_date__"] = labels[crawl_header]
    missing: list[str] = []
    for slot, label in layout["slots"].items():
        if label not in labels:
            missing.append(f"{slot} ({label!r})")
            continue
        mapping[slot] = labels[label]
    if missing:
        raise ManifestError(
            f"nguon {source_code!r}: workbook thieu cot slot {missing}. Header doc duoc: {list(labels)}."
        )
    return mapping


def _require_not_unevaluated_formula(source_code: str, cell: Any, slot: str, excel_row: int) -> None:
    raw = getattr(cell, "value", None)
    if isinstance(raw, str) and raw.startswith("="):
        raise ManifestError(
            f"nguon {source_code!r} dong {excel_row} slot {slot}: o chua CONG THUC {raw[:40]!r} nhung "
            f"KHONG co cached value - openpyxl khong tinh cong thuc. Mo workbook bang Excel/"
            f"LibreOffice, luu lai de cache gia tri, roi sinh lai manifest. KHONG implement lai "
            f"SLOT_RULES trong ETL."
        )


def _as_date(value: Any) -> dt.date | None:
    if value is None:
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        try:
            return dt.date.fromisoformat(text[:10])
        except ValueError:
            return None
    return None


def build_ownership_rows(workbooks: Mapping[str, str | Path]) -> list[OwnershipRow]:
    """Gop N workbook -> 1 bang phang. FAIL CUNG neu 2 nguon cung claim 1 (crawl_date, checkin)."""
    rows: list[OwnershipRow] = []
    for source_code in sorted(workbooks):
        rows.extend(read_plan_rows(source_code, workbooks[source_code]))
    assert_no_conflict(rows)
    rows.sort(key=lambda row: (row.crawl_date, row.checkin_date, row.owner_source, row.schedule_slot))
    return rows


def validate_rows(rows: Iterable[OwnershipRow]) -> None:
    """Fail-closed cho tung dong (GPT review 06 MAJOR 1).

    Ly do quan trong: mot typo `owner_source='vpz'` kem hash duoc sinh lai van load thanh cong o ban
    truoc, roi TOAN BO item cua VPS dung ngay se bi phan loai `non_owner_duplicate` va bi loai khoi
    main/reference/training - mot silent selection error, khong phai loi trinh bay.
    """
    problems: list[str] = []
    for index, row in enumerate(rows):
        layout = WORKBOOK_LAYOUTS.get(row.owner_source)
        if layout is None:
            problems.append(
                f"rows[{index}]: owner_source={row.owner_source!r} khong nam trong WORKBOOK_LAYOUTS "
                f"({sorted(WORKBOOK_LAYOUTS)}). Them nguon moi = khai bao layout truoc, khong nhan "
                f"chuoi tuy y."
            )
            continue
        slot = row.schedule_slot
        if not isinstance(slot, str) or not slot.strip():
            problems.append(f"rows[{index}]: schedule_slot rong.")
            continue
        if len(slot) > MAX_SLOT_LENGTH:
            problems.append(
                f"rows[{index}]: schedule_slot={slot!r} dai {len(slot)} ky tu, vuot "
                f"VARCHAR({MAX_SLOT_LENGTH}) cua etl_item_map."
            )
        if slot not in layout["slots"]:
            problems.append(
                f"rows[{index}]: schedule_slot={slot!r} khong thuoc slot cua nguon "
                f"{row.owner_source!r} ({sorted(layout['slots'])})."
            )
        if not isinstance(row.crawl_date, dt.date) or not isinstance(row.checkin_date, dt.date):
            problems.append(f"rows[{index}]: crawl_date/checkin_date khong phai kieu date.")
    if problems:
        raise ManifestError(
            f"ownership manifest co {len(problems)} dong khong hop le: " + "; ".join(problems[:10])
        )


def assert_no_conflict(rows: Iterable[OwnershipRow]) -> None:
    """Muc 8: 'Conflicting ownership: FAIL cung luc sinh manifest.'"""
    seen: dict[tuple[dt.date, dt.date], OwnershipRow] = {}
    conflicts: list[tuple[OwnershipRow, OwnershipRow]] = []
    for row in rows:
        key = (row.crawl_date, row.checkin_date)
        previous = seen.get(key)
        if previous is None:
            seen[key] = row
            continue
        conflicts.append((previous, row))
    if conflicts:
        detail = "; ".join(
            f"({a.crawl_date} / {a.checkin_date}) bi claim boi {a.owner_source}:{a.schedule_slot} "
            f"VA {b.owner_source}:{b.schedule_slot}"
            for a, b in conflicts[:10]
        )
        raise OwnershipConflictError(
            f"Conflicting ownership: {len(conflicts)} cap (crawl_date, checkin_date) co >1 chu so huu. "
            f"{detail}. Muc 8 yeu cau FAIL cung o buoc sinh manifest - nguoi van hanh phai quyet dinh "
            f"ai la chu so huu truoc khi build warehouse."
        )


def ownership_manifest_payload(rows: Sequence[OwnershipRow]) -> list[dict[str, Any]]:
    """LUON sort - hash phai la fingerprint cua TAP dong, khong phu thuoc thu tu ghi ra file.

    Neu khong sort o day: doi thu tu `--workbook` tren CLI se ra `ownership_manifest_sha256` khac,
    ma hash do lai la input cua MOI `schedule_manifest_row_key`/`schedule_day_key` - toan bo key se
    doi du lich khong he thay doi. Cung ly do voi `source_manifest_payload()`.
    """
    return sorted(
        (row.as_payload() for row in rows),
        key=lambda payload: (
            payload["crawl_date"], payload["checkin_date"],
            payload["owner_source"], payload["schedule_slot"],
        ),
    )


def compute_ownership_manifest_sha256(rows: Sequence[OwnershipRow]) -> str:
    """Hash tinh tu NOI DUNG cac dong, KHONG tu tham chieu envelope (yeu cau GPT file 02)."""
    return sha256_hex(canonical_json(ownership_manifest_payload(rows)))


# ======================================================================================
# Doc lai manifest + tra cuu
# ======================================================================================
@dataclass(frozen=True)
class OwnershipManifest:
    path: Path | None
    rows: tuple[OwnershipRow, ...]
    manifest_sha256: str
    _by_pair: dict[tuple[dt.date, dt.date], OwnershipRow]
    _source_crawl_dates: dict[str, frozenset[dt.date]]
    _source_checkins: dict[str, frozenset[dt.date]]
    _source_window: dict[str, tuple[dt.date, dt.date]]

    @classmethod
    def from_rows(cls, rows: Sequence[OwnershipRow], *, path: Path | None = None) -> "OwnershipManifest":
        validate_rows(rows)
        assert_no_conflict(rows)
        by_pair = {(row.crawl_date, row.checkin_date): row for row in rows}
        crawl_dates: dict[str, set[dt.date]] = {}
        checkins: dict[str, set[dt.date]] = {}
        for row in rows:
            crawl_dates.setdefault(row.owner_source, set()).add(row.crawl_date)
            checkins.setdefault(row.owner_source, set()).add(row.checkin_date)
        return cls(
            path=path,
            rows=tuple(rows),
            manifest_sha256=compute_ownership_manifest_sha256(list(rows)),
            _by_pair=by_pair,
            _source_crawl_dates={k: frozenset(v) for k, v in crawl_dates.items()},
            _source_checkins={k: frozenset(v) for k, v in checkins.items()},
            _source_window={k: (min(v), max(v)) for k, v in crawl_dates.items()},
        )

    def lookup(self, crawl_date: dt.date, checkin_date: dt.date) -> OwnershipRow | None:
        return self._by_pair.get((crawl_date, checkin_date))

    def source_plans_day(self, source_code: str, crawl_date: dt.date) -> bool:
        return crawl_date in self._source_crawl_dates.get(source_code, frozenset())

    def source_plans_checkin(self, source_code: str, checkin_date: dt.date) -> bool:
        return checkin_date in self._source_checkins.get(source_code, frozenset())

    def planned_window(self, source_code: str) -> tuple[dt.date, dt.date] | None:
        """(ngay crawl som nhat, muon nhat) co ke hoach cua nguon - None neu nguon khong co dong nao."""
        return self._source_window.get(source_code)

    def day_missing_reason(self, source_code: str, crawl_date: dt.date) -> str:
        """Phan biet 'truoc khi protocol bat dau' voi 'sau cua so' va 'trong cua so nhung trong lich'.

        GPT review 06 MINOR 1: truoc day moi ngay khong co plan deu bi goi la `pre_protocol_pilot`,
        nen mot run sau ngay cuoi workbook hoac mot source la cung bi ghi nhan sai trong audit.
        """
        window = self.planned_window(source_code)
        if window is None:
            return REASON_RUN_NOT_IN_MANIFEST
        first, last = window
        if crawl_date < first:
            return REASON_PRE_PROTOCOL_PILOT
        if crawl_date > last:
            return REASON_POST_PROTOCOL_WINDOW
        return REASON_DAY_NOT_IN_MANIFEST


def write_ownership_manifest(rows: Sequence[OwnershipRow], path: str | Path) -> str:
    """Ghi manifest JSON ATOMIC. Tra ve `ownership_manifest_sha256` da tinh.

    Atomic vi day la "EXTERNAL IMMUTABLE input" cua build: crash giua chung ma de lai 1 file JSON
    cut ngay tai duong dan do la truong hop te nhat (GPT review 06 MINOR 2). Cung convention voi
    promotion `warehouse_current.json`: temp file CUNG THU MUC -> flush+fsync -> os.replace.
    """
    import json
    import os
    import tempfile

    validate_rows(rows)
    payload = ownership_manifest_payload(rows)
    digest = sha256_hex(canonical_json(payload))
    document = {
        "manifest_version": 1,
        "ownership_manifest_sha256": digest,
        "row_count": len(payload),
        "sources": sorted({row.owner_source for row in rows}),
        "rows": payload,
    }
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    handle_fd, temp_name = tempfile.mkstemp(
        dir=str(target.parent), prefix=f".{target.name}.", suffix=".tmp"
    )
    try:
        with os.fdopen(handle_fd, "w", encoding="utf-8") as handle:
            json.dump(document, handle, ensure_ascii=False, indent=2)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temp_name, target)
    except BaseException:
        try:
            os.unlink(temp_name)
        except OSError:
            pass
        raise
    return digest


def load_ownership_manifest(path: str | Path) -> OwnershipManifest:
    import json

    manifest_path = Path(path)
    try:
        document = json.loads(manifest_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ManifestError(f"khong doc duoc ownership manifest {manifest_path}: {exc}") from exc
    if document.get("manifest_version") != 1:
        raise ManifestError(f"ownership manifest_version khong ho tro: {document.get('manifest_version')!r}")
    raw_rows = document.get("rows")
    if not isinstance(raw_rows, list) or not raw_rows:
        raise ManifestError("ownership manifest khong co 'rows' hoac rong.")

    rows: list[OwnershipRow] = []
    for index, raw in enumerate(raw_rows):
        try:
            rows.append(OwnershipRow(
                owner_source=raw["owner_source"],
                crawl_date=dt.date.fromisoformat(raw["crawl_date"]),
                schedule_slot=raw["schedule_slot"],
                checkin_date=dt.date.fromisoformat(raw["checkin_date"]),
            ))
        except (KeyError, TypeError, ValueError) as exc:
            raise ManifestError(f"ownership manifest rows[{index}] khong hop le: {exc}") from exc

    manifest = OwnershipManifest.from_rows(rows, path=manifest_path.resolve())

    declared_count = document.get("row_count")
    if declared_count != len(rows):
        raise ManifestError(
            f"envelope row_count={declared_count!r} khac so dong thuc te ({len(rows)}) - file khong nhat quan."
        )
    declared_sources = document.get("sources")
    actual_sources = sorted({row.owner_source for row in rows})
    if declared_sources != actual_sources:
        raise ManifestError(
            f"envelope sources={declared_sources!r} khac tap owner_source thuc te ({actual_sources})."
        )
    declared = document.get("ownership_manifest_sha256")
    if declared != manifest.manifest_sha256:
        raise ManifestError(
            f"ownership_manifest_sha256 trong file ({declared}) khac voi hash tinh lai tu chinh rows "
            f"({manifest.manifest_sha256}) - file da bi sua tay sau khi sinh."
        )
    return manifest


# ======================================================================================
# Cong thuc hash muc 8
# ======================================================================================
def schedule_manifest_row_key(
    *, ownership_manifest_sha256: str, owner_source: str, crawl_date: dt.date,
    schedule_slot: str, checkin_date: dt.date, cohort_manifest_sha256: str,
) -> str:
    """CHAR(64) cua `etl_item_map.schedule_manifest_row_key` - dinh danh 1 DONG CO THAT trong manifest."""
    return sha256_hex(canonical_json({
        "ownership_manifest_sha256": ownership_manifest_sha256,
        "owner_source": owner_source,
        "crawl_date": crawl_date.isoformat(),
        "schedule_slot": schedule_slot,
        "checkin_date": checkin_date.isoformat(),
        "cohort_manifest_sha256": cohort_manifest_sha256,
    }))


def schedule_day_key(
    *, ownership_manifest_sha256: str, owner_source: str, crawl_date: dt.date,
    cohort_manifest_sha256: str,
) -> str:
    """VARCHAR(100) cua `etl_run_map.schedule_day_key` - dinh danh 1 NGAY co ke hoach cua 1 nguon."""
    return sha256_hex(canonical_json({
        "ownership_manifest_sha256": ownership_manifest_sha256,
        "owner_source": owner_source,
        "crawl_date": crawl_date.isoformat(),
        "cohort_manifest_sha256": cohort_manifest_sha256,
    }))


# ======================================================================================
# Resolve (pure) - trai tim cua muc 8, test duoc khong can MySQL
# ======================================================================================
def resolve_run_ownership(
    *, manifest: OwnershipManifest, source_code: str, crawl_date: dt.date,
) -> RunOwnership:
    """Run-level: nguon nay co dong ke hoach nao cho DUNG ngay nay khong?"""
    if manifest.source_plans_day(source_code, crawl_date):
        return RunOwnership(
            planned_crawl_date=crawl_date, has_plan=True, flags=FLAGS_FULL, exclusion_reason=None,
        )
    # `chk_run_map_schedule` ep: khong co planned_crawl_date/schedule_day_key => 3 co phai FALSE.
    return RunOwnership(
        planned_crawl_date=None, has_plan=False, flags=FLAGS_RAW_ONLY,
        exclusion_reason=REASON_RUN_NOT_IN_MANIFEST,
    )


def resolve_item_ownership(
    *, manifest: OwnershipManifest, source_code: str, crawl_date: dt.date, checkin_date: dt.date,
    item_status: str, hotel_id: str | None, hotel_in_cohort: bool, run_flags: EligibilityFlags,
) -> ItemOwnership:
    """Item-level, theo dung thu tu da chot voi GPT (file 02 D2 / 02b / 04).

    Tra cuu dong manifest TRUOC, so owner SAU. `hotel_in_cohort` chi co y nghia khi
    `hotel_id is not None`; item `error` thuong co `hotel_id IS NULL` (local 478, vps 174 dong that)
    va khong duoc suy ra la ngoai cohort.
    """
    row = manifest.lookup(crawl_date, checkin_date)

    if row is None:
        if not manifest.source_plans_day(source_code, crawl_date):
            reason = manifest.day_missing_reason(source_code, crawl_date)
        elif manifest.source_plans_checkin(source_code, checkin_date):
            reason = REASON_OFF_PLAN_CHECKIN_OTHER_DAY
        else:
            reason = REASON_OFF_PLAN_UNKNOWN
        return ItemOwnership(
            ownership_status="unassigned", schedule_slot=None, owner_source=None,
            flags=FLAGS_RAW_ONLY.intersect(run_flags), exclusion_reason=reason,
        )

    if row.owner_source != source_code:
        return ItemOwnership(
            ownership_status="non_owner_duplicate", schedule_slot=row.schedule_slot,
            owner_source=row.owner_source, flags=FLAGS_RAW_ONLY.intersect(run_flags),
            exclusion_reason=REASON_NON_OWNER_DUPLICATE,
        )

    if hotel_id is not None and not hotel_in_cohort:
        return ItemOwnership(
            ownership_status="protocol_deviation", schedule_slot=row.schedule_slot,
            owner_source=row.owner_source, flags=FLAGS_RAW_ONLY.intersect(run_flags),
            exclusion_reason=REASON_HOTEL_OUTSIDE_COHORT,
        )

    if item_status == "success":
        return ItemOwnership(
            ownership_status="owner_success", schedule_slot=row.schedule_slot,
            owner_source=row.owner_source, flags=FLAGS_FULL.intersect(run_flags),
            exclusion_reason=None,
        )

    # GPT file 04 (O1/O2): moi status terminal khac tren owned slot deu la owner_failure; giu EDA
    # main de muc 13 van thay du success/partial/sold_out/not_bookable/error, nhung tat
    # reference/training de contract "owner failure khong fallback sang gia cua non-owner" doc lap
    # voi viec mot dataset_version tuong lai co the noi cho phep 'partial'.
    return ItemOwnership(
        ownership_status="owner_failure", schedule_slot=row.schedule_slot,
        owner_source=row.owner_source, flags=FLAGS_RAW_AND_MAIN_ONLY.intersect(run_flags),
        exclusion_reason=f"owner_failure_status_{item_status}",
    )
