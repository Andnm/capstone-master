"""Cohort manifest - nguon su that DUY NHAT cho `hotels.city` va tu cach thanh vien cohort (muc 6 rule 5).

Muc 6: "`city`/cohort **luon luon** lay tu cohort manifest, khong bao gio tu `hotels.city` nguon
nao". Ly do: moi nguon tu suy city tu address/sheet cua rieng no, 2 nguon co the bat dong; con
cohort manifest la file nguoi van hanh kiem soat, dung de dinh nghia quan the nghien cuu.

File nguon: `link_hotel_data_expanded.xlsx` o goc repo - 5 sheet, ten sheet CHINH LA ten thanh pho
(dung y het gia tri can co trong `hotels.city`), moi sheet 2 cot: ten khach san + link Booking.
`hotel_id` suy tu link bang `extract_hotel_slug()` cua chinh scraper - KHONG viet lai logic slug.

`cohort_manifest_sha256` tinh tu NOI DUNG da chuan hoa (danh sach (hotel_id, city) sap xep), khong
phai tu bytes cua file .xlsx - file Excel doi bytes moi lan mo/luu du noi dung khong doi.

COHORT THEO VERSION (phat hien khi rehearsal batch 2): cohort KHONG phai 1 workbook duy nhat ma la lich su
version co ngay hieu luc - CLAUDE.md muc 2: v1=355 dung cho moi run truoc 02/09, v2=354 tu do (Mac Valley
`mac-dalat` bi go khoi Booking). Workbook o goc repo bi SUA TAI CHO, nen tra cohort bang ban hien tai se
danh 108 item success hop le 18-26/08 cua Mac Valley la `protocol_deviation` va city=NULL -> loai khoi
train, dung loai survivor-selection bias ma CLAUDE.md cam. `CohortHistory`: thanh vien cua 1 item =
version CO HIEU LUC tai crawl_date VN cua run; city = hop cac version (1 hotel chi duoc co 1 city).
"""
from __future__ import annotations

import datetime as dt
import json
from dataclasses import dataclass
from pathlib import Path
from types import MappingProxyType
from typing import Iterable, Mapping

from app.scraper.url_utils import extract_hotel_slug

from .errors import ManifestError
from .hashing import canonical_json, sha256_hex

# CLAUDE.md muc 2/7.2 - 5 thanh pho trong scope, viet DUNG nhu vay de JOIN thang voi hotels.city.
VALID_CITIES = ("Hồ Chí Minh", "Hà Nội", "Vũng Tàu", "Đà Lạt", "Phú Quốc")
COHORT_HISTORY_VERSION = 1

_LINK_HEADER_CANDIDATES = ("link", "url")
_HISTORY_REQUIRED = ("cohort_version", "effective_from_crawl_date", "workbook_path", "members_sha256", "size")


@dataclass(frozen=True)
class CohortManifest:
    """`hotel_city` la MappingProxyType - `frozen=True` chi chan gan lai field, KHONG chan sua dict
    ben trong. Neu de dict mutable thi `manifest_sha256` va `city_of()` co the lech nhau sau khi
    load (GPT review 06 MINOR 4)."""

    path: Path
    hotel_city: Mapping[str, str]

    @property
    def manifest_sha256(self) -> str:
        return sha256_hex(canonical_json(sorted(self.hotel_city.items())))

    @property
    def size(self) -> int:
        return len(self.hotel_city)

    def city_of(self, hotel_id: str) -> str | None:
        """None = hotel KHONG thuoc cohort -> warehouse ghi city=NULL (muc 6 rule 5, chot o file 02)."""
        return self.hotel_city.get(hotel_id)

    def contains(self, hotel_id: str) -> bool:
        return hotel_id in self.hotel_city


def load_cohort_manifest(path: str | Path, *, require_all_cities: bool = True) -> CohortManifest:
    """Doc workbook cohort -> {hotel_id: city}. FAIL CLOSED tren moi bat thuong.

    `require_all_cities=True` (mac dinh cho moi duong chay that): thieu 1 sheet thanh pho la FAIL,
    khong phai canh bao. Neu khong, xoa nham 1 sheet se lam scope dataset AM THAM co lai tu 5 thanh
    pho xuong 4 ma khong ai phat hien (GPT review 06 MINOR 4). Chi test fixture moi truyen False.
    """
    import openpyxl  # import cuc bo: chi lenh warehouse can openpyxl, worker crawl thi khong

    workbook_path = Path(path)
    if not workbook_path.exists():
        raise ManifestError(f"khong tim thay cohort manifest: {workbook_path}")

    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        unknown_sheets = [name for name in workbook.sheetnames if name not in VALID_CITIES]
        if unknown_sheets:
            raise ManifestError(
                f"cohort manifest co sheet khong phai thanh pho trong scope: {unknown_sheets}. "
                f"Chi chap nhan dung {list(VALID_CITIES)} (CLAUDE.md muc 2)."
            )
        missing_cities = [city for city in VALID_CITIES if city not in workbook.sheetnames]
        if require_all_cities and missing_cities:
            raise ManifestError(
                f"cohort manifest THIEU sheet thanh pho: {missing_cities}. Scope de tai co dinh 5 "
                f"thanh pho (CLAUDE.md muc 2) - thieu sheet se lam scope dataset am tham co lai."
            )

        hotel_city: dict[str, str] = {}
        duplicates: list[tuple[str, str, str]] = []
        for city in workbook.sheetnames:
            sheet = workbook[city]
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, None)
            link_index = _find_link_column(header, city)
            for row_number, row in enumerate(rows, start=2):
                if row is None or all(cell is None for cell in row):
                    continue
                raw_link = row[link_index] if link_index < len(row) else None
                if raw_link is None or not str(raw_link).strip():
                    raise ManifestError(
                        f"cohort manifest sheet {city!r} dong {row_number}: thieu link Booking - "
                        f"khong suy duoc hotel_id."
                    )
                hotel_id = extract_hotel_slug(str(raw_link))
                if not hotel_id:
                    raise ManifestError(
                        f"cohort manifest sheet {city!r} dong {row_number}: khong parse duoc slug tu "
                        f"link {str(raw_link)[:80]!r}."
                    )
                if hotel_id in hotel_city:
                    duplicates.append((hotel_id, hotel_city[hotel_id], city))
                    continue
                hotel_city[hotel_id] = city
    finally:
        workbook.close()

    if duplicates:
        raise ManifestError(
            f"cohort manifest co hotel_id xuat hien o nhieu sheet (khong the xac dinh city): "
            f"{duplicates[:10]}"
        )
    if not hotel_city:
        raise ManifestError("cohort manifest rong - khong doc duoc hotel nao.")
    return CohortManifest(
        path=workbook_path.resolve(), hotel_city=MappingProxyType(dict(hotel_city))
    )


def _find_link_column(header: tuple | None, city: str) -> int:
    if not header:
        raise ManifestError(f"cohort manifest sheet {city!r} khong co dong header.")
    for index, cell in enumerate(header):
        if cell is None:
            continue
        text = str(cell).strip().lower()
        if any(candidate in text for candidate in _LINK_HEADER_CANDIDATES):
            return index
    raise ManifestError(
        f"cohort manifest sheet {city!r}: khong tim thay cot link trong header {header!r} "
        f"(tim theo tu khoa {_LINK_HEADER_CANDIDATES})."
    )


# ======================================================================== cohort theo version
@dataclass(frozen=True)
class CohortVersion:
    label: str
    effective_from: dt.date  # crawl_date VN dau tien version nay co hieu luc
    manifest: CohortManifest
    git_commit: str | None = None
    reason: str | None = None

    def summary(self) -> dict:
        return {"cohort_version": self.label, "effective_from_crawl_date": self.effective_from.isoformat(),
                "size": self.manifest.size, "members_sha256": self.manifest.manifest_sha256,
                "git_commit": self.git_commit}


@dataclass(frozen=True)
class CohortHistory:
    """Chi tao qua `cohort_history_from_versions` (validate thu tu, nhan, city)."""

    versions: tuple[CohortVersion, ...]
    hotel_city: Mapping[str, str]
    path: Path | None = None

    @property
    def manifest_sha256(self) -> str:
        """Hash NOI DUNG ca lich su: doi ngay hieu luc hoac thanh vien cua 1 version -> hash doi."""
        return sha256_hex(canonical_json([
            {"cohort_version": version.label, "effective_from_crawl_date": version.effective_from.isoformat(),
             "members_sha256": version.manifest.manifest_sha256}
            for version in self.versions
        ]))

    def version_at(self, crawl_date: dt.date) -> CohortVersion | None:
        effective = None
        for version in self.versions:  # da ep tang dan nghiem ngat
            if version.effective_from <= crawl_date:
                effective = version
        return effective

    def contains_at(self, hotel_id: str, crawl_date: dt.date) -> bool:
        """Thanh vien theo version CO HIEU LUC tai crawl_date. Truoc version dau tien -> khong ai."""
        version = self.version_at(crawl_date)
        return version is not None and version.manifest.contains(hotel_id)

    def city_of(self, hotel_id: str) -> str | None:
        """Hop moi version: hotel da roi cohort van giu city (du lieu truoc khi roi van dung duoc)."""
        return self.hotel_city.get(hotel_id)

    def summary(self) -> list[dict]:
        return [version.summary() for version in self.versions]


def cohort_history_from_versions(versions: Iterable[CohortVersion], *, path: Path | None = None) -> CohortHistory:
    ordered = tuple(versions)
    if not ordered:
        raise ManifestError("cohort history rong - can it nhat 1 version.")
    labels = [version.label for version in ordered]
    if len(set(labels)) != len(labels):
        raise ManifestError(f"cohort history trung nhan version: {labels}")
    for earlier, later in zip(ordered, ordered[1:]):
        if later.effective_from <= earlier.effective_from:
            raise ManifestError(
                f"effective_from_crawl_date phai tang dan nghiem ngat: {earlier.label}={earlier.effective_from} "
                f"-> {later.label}={later.effective_from}."
            )
    hotel_city: dict[str, str] = {}
    conflicts: list[tuple[str, str, str, str]] = []
    for version in ordered:
        for hotel_id, city in version.manifest.hotel_city.items():
            known = hotel_city.setdefault(hotel_id, city)
            if known != city:
                conflicts.append((hotel_id, known, city, version.label))
    if conflicts:
        raise ManifestError(
            f"1 hotel co city KHAC NHAU giua cac version cohort: {conflicts[:10]} - city phai co dinh theo "
            f"hotel_id, khong the chon ngam 1 ben."
        )
    return CohortHistory(versions=ordered, hotel_city=MappingProxyType(hotel_city), path=path)


def single_version_history(manifest: CohortManifest) -> CohortHistory:
    """1 workbook hieu luc MOI ngay. CHI cho fixture/test: build that phai dung history JSON (CLI
    `build_warehouse.py` tu choi .xlsx) vi workbook hien tai da mat cac hotel roi cohort."""
    return cohort_history_from_versions([CohortVersion("single", dt.date.min, manifest)], path=manifest.path)


def load_cohort_history(path: str | Path, *, base_dir: str | Path) -> CohortHistory:
    """Doc cohort history JSON. FAIL CLOSED: moi workbook phai khop `members_sha256` + `size` da khai bao
    (workbook bi sua sau khi tao history -> dung ngay, khong am tham dung noi dung moi)."""
    history_path = Path(path)
    if not history_path.exists():
        raise ManifestError(f"khong tim thay cohort history: {history_path}")
    try:
        payload = json.loads(history_path.read_text(encoding="utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ManifestError(f"cohort history {history_path.name} khong phai JSON hop le: {exc}") from exc
    if not isinstance(payload, dict) or payload.get("cohort_history_version") != COHORT_HISTORY_VERSION:
        raise ManifestError(
            f"cohort history {history_path.name}: cohort_history_version phai = {COHORT_HISTORY_VERSION}."
        )
    entries = payload.get("versions")
    if not isinstance(entries, list) or not entries:
        raise ManifestError(f"cohort history {history_path.name}: 'versions' phai la list khong rong.")
    versions: list[CohortVersion] = []
    for index, entry in enumerate(entries):
        missing = [key for key in _HISTORY_REQUIRED if not isinstance(entry, dict) or key not in entry]
        if missing:
            raise ManifestError(f"cohort history version #{index}: thieu truong {missing}.")
        label = str(entry["cohort_version"])
        try:
            effective = dt.date.fromisoformat(str(entry["effective_from_crawl_date"]))
        except ValueError as exc:
            raise ManifestError(f"cohort history version {label}: effective_from_crawl_date sai dinh dang.") from exc
        workbook = Path(entry["workbook_path"])
        if not workbook.is_absolute():
            workbook = Path(base_dir) / workbook
        manifest = load_cohort_manifest(workbook)
        if manifest.manifest_sha256 != entry["members_sha256"]:
            raise ManifestError(
                f"cohort history version {label}: members_sha256 KHONG KHOP - workbook {workbook} da doi noi "
                f"dung sau khi tao history (khai bao {entry['members_sha256']}, doc duoc {manifest.manifest_sha256})."
            )
        if type(entry["size"]) is not int or manifest.size != entry["size"]:
            raise ManifestError(
                f"cohort history version {label}: size khai bao {entry['size']!r} != so hotel doc duoc {manifest.size}."
            )
        versions.append(CohortVersion(label, effective, manifest, entry.get("git_commit"), entry.get("reason")))
    return cohort_history_from_versions(versions, path=history_path.resolve())


def load_cohort(path: str | Path, *, base_dir: str | Path) -> CohortHistory:
    """`.json` -> cohort history nhieu version (duong chay that); con lai -> 1 workbook (fixture)."""
    if Path(path).suffix.lower() == ".json":
        return load_cohort_history(path, base_dir=base_dir)
    return single_version_history(load_cohort_manifest(path))
